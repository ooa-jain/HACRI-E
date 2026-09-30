"""
The AI survey (outcome & impact) one department at a time: before and after,
section by section, PRaiSE and entrepreneur families — on a page, behind a
share link, and in a workbook.

Six students, by hand:

  Law (Bangalore)  A, B filled both; C filled only the baseline
  Law (Kochi)      F filled both
  Commerce         D filled only the baseline; E filled nothing

Every baseline scores 2 out of 5 on every item and every post survey 4, so
each average is 2.0 before and 4.0 after. PRaiSE: A Entrepreneurship,
B Social Good, F None. A's father and F's mother are entrepreneurs.
"""
from __future__ import annotations

import io
from datetime import datetime, timezone
from typing import AsyncIterator
from urllib.parse import urlsplit

import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from mongomock_motor import AsyncMongoMockClient

from app import db


@pytest_asyncio.fixture
async def app_with_mock():
    db._set_client_for_tests(AsyncMongoMockClient())
    try:
        from app.main import app
        await db.init_indexes(allow_duplicate_email=True)
        yield app
    finally:
        db._reset_clients_for_tests()


@pytest_asyncio.fixture
async def admin_client(app_with_mock) -> AsyncIterator[AsyncClient]:
    async with AsyncClient(transport=ASGITransport(app=app_with_mock), base_url="http://test") as ac:
        ac.cookies.set("survey_admin_session", "1")
        yield ac


@pytest_asyncio.fixture
async def client(app_with_mock) -> AsyncIterator[AsyncClient]:
    async with AsyncClient(transport=ASGITransport(app=app_with_mock), base_url="http://test") as ac:
        yield ac


def _sheet(value: int) -> dict:
    from app.hacri_e2_compat import SCHEMA
    return {key: str(6 - value if rev else value) for key, (_d, rev) in SCHEMA.items()}


POST_EXTRA = {
    "a@x.com": {"praise_initiative": "Entrepreneurship", "father_occupation": "Entrepreneur",
                "father_name": "Mr A", "business_name": "Acme Traders", "business_type": "Retail"},
    "b@x.com": {"praise_initiative": "Social Good", "father_occupation": "Salaried"},
    "f@x.com": {"praise_initiative": "None", "mother_occupation": "Entrepreneur",
                "mother_name": "Mrs F", "mother_business_name": "F Bakes",
                "mother_business_type": "Food"},
}


async def _seed() -> None:
    now = datetime.now(timezone.utc)
    people = [
        ("a@x.com", "Department of Law", "Bangalore", True, True),
        ("b@x.com", "Department of Law", "Bangalore", True, True),
        ("c@x.com", "Department of Law", "Bangalore", True, False),
        ("d@x.com", "Department of Commerce", "Bangalore", True, False),
        ("e@x.com", "Department of Commerce", "Bangalore", False, False),
        ("f@x.com", "Department of Law", "Kochi", True, True),
    ]
    for email, dept, campus, pre, post in people:
        status = db.STATUS_POST_DONE if post else db.STATUS_PRE_DONE if pre else None
        await db.get_db()["users"].insert_one({
            "email": email, "name": email[0].upper(), "program": dept, "ug_or_pg": "ug",
            "location": campus, "status": status, "created_at": now,
        })
        if pre:
            await db.get_db()["pre_responses"].insert_one(
                {"email": email, "submitted_at": now, "fields": _sheet(2)})
        if post:
            await db.get_db()["post_responses"].insert_one(
                {"email": email, "submitted_at": now,
                 "fields": {**_sheet(4), **POST_EXTRA.get(email, {})}})


def _law(pack):
    return next(d for d in pack["departments"] if d["dept"] == "Department of Law")


# ── The analysis ─────────────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_a_department_reads_before_after_and_how_far_it_rose(app_with_mock):
    await _seed()
    from app.outcome_departments import outcome_pack
    pack = await outcome_pack()
    law = _law(pack)

    assert law["journey"]["registered"] == 4
    assert law["journey"]["baseline"] == 4 and law["journey"]["post"] == 3
    assert law["journey"]["pending_post"] == 1
    for key in ("literacy", "readiness", "overall"):
        assert (law["scores"][key]["before"], law["scores"][key]["after"],
                law["scores"][key]["change"]) == (2.0, 4.0, 2.0)
    m = law["movement"]
    assert (m["matched"], m["gained"], m["unchanged"], m["declined"]) == (3, 3, 0, 0)
    # Section by section, reversed items flipped so every section reads 2 → 4.
    assert [s["key"] for s in law["sections"]] == ["B", "D", "E", "F", "G"]
    for s in law["sections"]:
        assert (s["before"], s["after"], s["change"]) == (2.0, 4.0, 2.0)
    assert "4 registered; 4 filled the baseline survey" in law["synopsis"]
    assert "rose from 2.0 to 4.0" in law["synopsis"]

    # The whole scope reads the same way, and every department is there.
    assert pack["overall"]["journey"]["registered"] == 6
    assert [d["dept"] for d in pack["departments"]] == ["Department of Law",
                                                         "Department of Commerce"]


@pytest.mark.asyncio
async def test_praise_and_entrepreneur_families_are_counted_and_named(app_with_mock):
    await _seed()
    from app.outcome_departments import outcome_pack
    law = _law(await outcome_pack())

    p = law["praise"]
    assert (p["answered"], p["joining"], p["none"]) == (3, 2, 1)
    pillars = {x["pillar"]: x for x in p["pillars"]}
    assert [s["name"] for s in pillars["Entrepreneurship"]["students"]] == ["A"]
    assert [s["name"] for s in pillars["Social Good"]["students"]] == ["B"]
    assert pillars["Human Excellence"]["count"] == 0

    e = law["entrepreneurs"]
    assert (e["count"], e["of_post"]) == (2, 3)
    by_name = {s["name"]: s["parents"] for s in e["students"]}
    assert by_name["A"] == [{"parent": "Father", "name": "Mr A",
                             "business": "Acme Traders", "type": "Retail"}]
    assert by_name["F"][0]["parent"] == "Mother" and by_name["F"][0]["business"] == "F Bakes"


@pytest.mark.asyncio
async def test_a_campus_narrows_the_department_to_that_campus(app_with_mock):
    await _seed()
    from app.outcome_departments import department_outcome_for
    kochi = await department_outcome_for("Department of Law", campus="Kochi")
    assert kochi["journey"]["registered"] == 1
    assert [s["name"] for s in kochi["students"]] == ["F"]


@pytest.mark.asyncio
async def test_the_quadrant_chart_follows_only_students_who_filled_both(app_with_mock):
    """C filled only the baseline: one end of a line is not a line, so C is
    not on the chart. A, B and F filled both."""
    await _seed()
    from app.outcome_departments import outcome_pack
    law = _law(await outcome_pack())
    chart = law["quadrant_chart"]

    assert chart["matched"] == 3
    assert sorted(p["name"] for p in chart["pairs"]) == ["A", "B", "F"]
    for p in chart["pairs"]:
        assert (p["pre_lit"], p["pre_read"], p["post_lit"], p["post_read"]) == (2.0, 2.0, 4.0, 4.0)
        assert p["rising"] and p["moved"]           # Novice → Champion
    assert (chart["moved"], chart["backwards"]) == (3, 0)
    # The averages are of these same students, and sit where the maths says.
    assert (chart["pre_centre"]["lit"], chart["pre_centre"]["read"]) == (2.0, 2.0)
    assert (chart["post_centre"]["lit"], chart["post_centre"]["read"]) == (4.0, 4.0)
    assert chart["post_centre"]["x"] > chart["geo"]["mid_x"] > chart["pre_centre"]["x"]
    assert chart["post_centre"]["y"] < chart["geo"]["mid_y"] < chart["pre_centre"]["y"]

    # Commerce has nobody who filled both — an empty chart, not a broken one.
    commerce = next(d for d in (await outcome_pack())["departments"]
                    if d["dept"] == "Department of Commerce")
    assert commerce["quadrant_chart"]["matched"] == 0
    assert commerce["quadrant_chart"]["pre_centre"] is None


# ── The pages ────────────────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_the_department_page_draws_the_literacy_readiness_chart(client, admin_client):
    await _seed()
    page = (await admin_client.get("/admin/survey/outcome/department",
                                   params={"dept": "Department of Law"})).text
    assert "the students who filled both surveys" in page
    assert page.count('marker-end="url(#oq-up)"') == 3          # one line per student
    assert "Q1  AI Champion" in page and "Q3  AI Novice" in page
    assert "Baseline average" in page and "Post average" in page
    assert "AI Literacy — do I understand AI?" in page
    assert "3 students filled both surveys." in page
    assert "3 moved to a different quadrant." in page

    empty = (await admin_client.get("/admin/survey/outcome/department",
                                    params={"dept": "Department of Commerce"})).text
    assert "Nobody in this department has filled both surveys yet" in empty
    assert "<svg viewBox" not in empty

    # A shared page draws it too, and still leaves emails out.
    from app.routes.shared_analysis import get_outcome_dept_token
    shared = (await client.get("/shared/outcome/department", params={
        "dept": "Department of Law",
        "token": get_outcome_dept_token("Department of Law")})).text
    assert shared.count('marker-end="url(#oq-up)"') == 3 and "@x.com" not in shared

@pytest.mark.asyncio
async def test_the_admin_department_page_shows_everything_with_emails(client, admin_client):
    await _seed()
    url = "/admin/survey/outcome/department"
    assert (await client.get(url, params={"dept": "Department of Law"})).status_code == 403
    assert (await admin_client.get(url, params={"dept": "Nope"})).status_code == 404

    page = (await admin_client.get(url, params={"dept": "Department of Law"})).text
    assert "<h1>Department of Law</h1>" in page
    for heading in ("Before and after the workshop", "Section by section",
                    "Where students sit", "Who wants to contribute to PRaiSE",
                    "Students from an entrepreneur family", "Every student"):
        assert heading in page
    assert "Acme Traders" in page and "a@x.com" in page
    assert "/admin/survey/outcome/department.xlsx?dept=Department+of+Law" in page
    assert "/shared/outcome/department?" in page


@pytest.mark.asyncio
async def test_a_share_link_opens_one_department_by_name_but_without_emails(client):
    await _seed()
    from app.routes.shared_analysis import get_cohort_token, get_outcome_dept_token
    url = "/shared/outcome/department"
    law = get_outcome_dept_token("Department of Law")

    r = await client.get(url, params={"dept": "Department of Law", "token": law})
    assert r.status_code == 200
    assert "<h1>Department of Law</h1>" in r.text and "Acme Traders" in r.text
    assert "@x.com" not in r.text
    # Not another department, not another campus, and not the cohort's own token.
    assert (await client.get(url, params={"dept": "Department of Commerce",
                                          "token": law})).status_code == 403
    assert (await client.get(url, params={"dept": "Department of Law", "campus": "Kochi",
                                          "token": law})).status_code == 403
    assert (await client.get(url, params={"dept": "Department of Law",
                                          "token": get_cohort_token()})).status_code == 403


# ── The workbook ─────────────────────────────────────────────────────────────

def _book(content: bytes):
    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(content))


def _headers(ws) -> list[str]:
    return [c.value for c in ws[4] if c.value is not None]


@pytest.mark.asyncio
async def test_the_all_departments_workbook_carries_every_analysis(client, admin_client):
    await _seed()
    assert (await client.get("/admin/survey/outcome/departments.xlsx")).status_code == 403
    r = await admin_client.get("/admin/survey/outcome/departments.xlsx")
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    wb = _book(r.content)
    assert wb.sheetnames == ["Summary", "Sections", "Quadrants", "PRaiSE",
                             "Entrepreneurs", "Students"]

    summary = wb["Summary"]
    names = [summary.cell(row=i, column=1).value for i in range(5, summary.max_row + 1)]
    assert names == ["All departments", "Department of Law", "Department of Commerce"]
    head = _headers(summary)
    law_row = [summary.cell(row=6, column=i + 1).value for i in range(len(head))]
    law = dict(zip(head, law_row))
    assert law["Baseline filled (before)"] == 4 and law["Post filled (after)"] == 3
    assert (law["Overall before"], law["Overall after"], law["Overall change"]) == (2.0, 4.0, 2.0)
    assert law["Want to join PRaiSE"] == 2 and law["Entrepreneur families"] == 2

    assert "Email" in _headers(wb["Students"])
    assert wb["Students"].max_row - 4 == 6          # every student
    assert wb["Entrepreneurs"].max_row - 4 == 2
    assert wb["PRaiSE"].max_row - 4 == 2            # None is counted, not listed


@pytest.mark.asyncio
async def test_a_department_workbook_holds_that_department_and_shared_ones_drop_emails(
        client, admin_client):
    await _seed()
    r = await admin_client.get("/admin/survey/outcome/department.xlsx",
                               params={"dept": "Department of Law"})
    wb = _book(r.content)
    summary = wb["Summary"]
    assert [summary.cell(row=i, column=1).value for i in range(5, summary.max_row + 1)] == [
        "Department of Law"]
    assert "Email" in _headers(wb["Students"])

    from app.routes.shared_analysis import get_outcome_dept_token
    shared = await client.get("/shared/outcome/department.xlsx",
                              params={"dept": "Department of Law",
                                      "token": get_outcome_dept_token("Department of Law")})
    assert shared.status_code == 200
    swb = _book(shared.content)
    for sheet in ("Students", "PRaiSE", "Entrepreneurs"):
        assert "Email" not in _headers(swb[sheet])
    assert (await client.get("/shared/outcome/department.xlsx",
                             params={"dept": "Department of Law",
                                     "token": "nope"})).status_code == 403


# ── The dashboard list ───────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_the_dashboard_lists_every_department_with_working_links(client, admin_client):
    await _seed()
    assert (await client.get("/admin/survey/outcome/department-links")).status_code == 403
    rows = (await admin_client.get("/admin/survey/outcome/department-links")).json()["departments"]
    assert [r["dept"] for r in rows] == ["Department of Law", "Department of Commerce"]
    law = rows[0]
    assert (law["registered"], law["baseline"], law["post"]) == (4, 4, 3)
    assert (law["before"], law["after"], law["change"]) == (2.0, 4.0, 2.0)
    assert (law["praise_joining"], law["entrepreneurs"]) == (2, 2)

    assert (await admin_client.get(law["open_url"])).status_code == 200
    assert (await admin_client.get(law["excel_url"])).status_code == 200
    share = urlsplit(law["share_url"])
    assert (await client.get(f"{share.path}?{share.query}")).status_code == 200

    from pathlib import Path
    tpl = Path("app/templates/admin_survey.html").read_text()
    assert "Department AI survey pages" in tpl
    assert "/admin/survey/outcome/department-links?campus=" in tpl
    assert "/admin/survey/outcome/departments.xlsx" in tpl


@pytest.mark.asyncio
async def test_the_shared_outcome_page_lists_every_department_with_its_own_links(client):
    await _seed()
    from app.routes.shared_analysis import get_cohort_token
    path = "/shared/cohort/departments"

    # Only the cohort link's own token lists them.
    assert (await client.get(path, params={"token": "nope"})).status_code == 403
    page = await client.get("/shared/cohort", params={"token": get_cohort_token()})
    assert page.status_code == 200 and "Every department's own analysis" in page.text
    assert "loadDeptPages()" in page.text

    r = await client.get(path, params={"token": get_cohort_token()})
    assert r.status_code == 200
    rows = r.json()["departments"]
    assert [x["dept"] for x in rows] == ["Department of Law", "Department of Commerce"]
    law = rows[0]
    assert (law["registered"], law["baseline"], law["post"]) == (4, 4, 3)
    assert (law["before"], law["after"], law["entrepreneurs"]) == (2.0, 4.0, 2)

    # Shared links only — no admin URLs, no emails — and each one opens its department.
    assert "/admin/" not in r.text and "@x.com" not in r.text
    for key, expect in (("open_url", "<h1>Department of Law</h1>"), ("excel_url", None)):
        u = urlsplit(law[key])
        got = await client.get(f"{u.path}?{u.query}")
        assert got.status_code == 200
        if expect:
            assert expect in got.text

    # A Kochi link lists only that campus's departments.
    kochi = await client.get(path, params={"token": get_cohort_token("Kochi"), "campus": "Kochi"})
    assert [x["dept"] for x in kochi.json()["departments"]] == ["Department of Law"]
    assert kochi.json()["departments"][0]["registered"] == 1


@pytest.mark.asyncio
async def test_each_department_page_links_to_every_other_department(client, admin_client):
    await _seed()
    from app.routes.shared_analysis import get_outcome_dept_token

    # Admin: admin links, biggest department first, this one marked.
    a = (await admin_client.get("/admin/survey/outcome/department",
                                params={"dept": "Department of Law"})).text
    assert 'aria-label="Departments"' in a
    assert a.index("Department of Law <b>4</b>") < a.index("Department of Commerce <b>2</b>")
    assert 'class="on" aria-current="page">Department of Law' in a
    assert "/admin/survey/outcome/department?dept=Department+of+Commerce" in a

    # Shared: shared links that open, each with its own token, and no emails.
    s = (await client.get("/shared/outcome/department", params={
        "dept": "Department of Law", "token": get_outcome_dept_token("Department of Law")})).text
    assert "/admin/" not in s.split('aria-label="Departments"')[1].split("</nav>")[0]
    tok = get_outcome_dept_token("Department of Commerce")
    assert f"dept=Department+of+Commerce&amp;token={tok}" in s or f"token={tok}" in s
    assert "@x.com" not in s
