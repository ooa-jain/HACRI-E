"""
outcome_excel.py — the AI survey's department analysis as a workbook.

Built from `outcome_departments.department_outcome()` results, so every figure
here is the one the department's own page shows. One workbook serves both the
whole cohort (every department, one row each) and a single department (the
same sheets, holding only that department).

Sheets: Summary · Sections · Quadrants · PRaiSE · Entrepreneurs · Students.
Share links pass `include_emails=False`, which drops every email column.
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.outcome_departments import PRAISE_PILLARS, SECTIONS
from app.sections import SECTION_TITLES

NAVY = "1E4F9A"
MIST = "F2F5FB"
HEAD_FILL = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
BAND_FILL = PatternFill(start_color=MIST, end_color=MIST, fill_type="solid")
HEAD_FONT = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Segoe UI", size=15, bold=True, color=NAVY)
NOTE_FONT = Font(name="Segoe UI", size=10, italic=True, color="5B6570")
BODY = Font(name="Segoe UI", size=11)
BOLD = Font(name="Segoe UI", size=11, bold=True)


def _sheet(wb, title: str, heading: str, note: str):
    ws = wb.create_sheet(title[:31])
    ws["A1"] = heading
    ws["A1"].font = TITLE_FONT
    ws["A2"] = note
    ws["A2"].font = NOTE_FONT
    ws.freeze_panes = "A5"
    return ws


def _table(ws, headers: list[str], rows: list[list], widths: list[int], *, start: int = 4,
           bold_first_row: bool = False):
    for col, (name, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=start, column=col, value=name)
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[start].height = 32
    for i, values in enumerate(rows, start=1):
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=start + i, column=col, value=value)
            cell.font = BOLD if (bold_first_row and i == 1) else BODY
            if i % 2 == 0:
                cell.fill = BAND_FILL
    if rows:
        ws.auto_filter.ref = f"A{start}:{get_column_letter(len(headers))}{start + len(rows)}"


def _drop(headers: list[str], rows: list[list], widths: list[int], name: str):
    """Remove one column by header name."""
    if name not in headers:
        return headers, rows, widths
    i = headers.index(name)
    return (headers[:i] + headers[i + 1:],
            [r[:i] + r[i + 1:] for r in rows],
            widths[:i] + widths[i + 1:])


def build_outcome_workbook(departments: list[dict], *, scope: str, generated_at: str,
                           overall: dict | None = None, include_emails: bool = True) -> bytes:
    """`departments` are department_outcome() results; `overall` leads the summary."""
    wb = Workbook()
    wb.remove(wb.active)
    every = ([overall] if overall else []) + departments

    # ── Summary ─────────────────────────────────────────────────────────────
    ws = _sheet(wb, "Summary", f"AI survey — outcome & impact by department · {scope}",
                f"Generated {generated_at}. Scores are 1–5 averages. Before = baseline survey, "
                "after = post survey. Change = after − before.")
    headers = ["Department", "Registered", "Baseline filled (before)", "Baseline %",
               "Post filled (after)", "Post %", "Still owe post",
               "Literacy before", "Literacy after", "Literacy change",
               "Readiness before", "Readiness after", "Readiness change",
               "Overall before", "Overall after", "Overall change",
               "Filled both", "Improved", "Held steady", "Slipped", "Improved %",
               "PRaiSE answered", "Want to join PRaiSE", "Join %",
               *[f"PRaiSE: {p}" for p in PRAISE_PILLARS], "PRaiSE: None",
               "Entrepreneur families", "Entrepreneur %"]
    rows = []
    for o in every:
        j, s, m, p, e = o["journey"], o["scores"], o["movement"], o["praise"], o["entrepreneurs"]
        rows.append([
            o["dept"], j["registered"], j["baseline"], j["baseline_pct"], j["post"], j["post_pct"],
            j["pending_post"],
            s["literacy"]["before"], s["literacy"]["after"], s["literacy"]["change"],
            s["readiness"]["before"], s["readiness"]["after"], s["readiness"]["change"],
            s["overall"]["before"], s["overall"]["after"], s["overall"]["change"],
            m["matched"], m["gained"], m["unchanged"], m["declined"], m["gained_pct"],
            p["answered"], p["joining"], p["joining_pct"],
            *[x["count"] for x in p["pillars"]], p["none"],
            e["count"], e["pct"],
        ])
    _table(ws, headers, rows, [42] + [13] * (len(headers) - 1), bold_first_row=bool(overall))

    # ── Sections ────────────────────────────────────────────────────────────
    ws = _sheet(wb, "Sections", "Section by section — before and after",
                "Each section's 1–5 average on the baseline and the post survey, reversed items "
                "flipped. G is the entrepreneurship section.")
    headers = ["Department", "Section", "Title", "Before", "After", "Change",
               "Answered before", "Answered after"]
    rows = [[o["dept"], s["key"], s["title"], s["before"], s["after"], s["change"],
             s["before_n"], s["after_n"]]
            for o in every for s in o["sections"]]
    _table(ws, headers, rows, [42, 9, 44, 11, 11, 11, 14, 14])

    # ── Quadrants and bands ─────────────────────────────────────────────────
    ws = _sheet(wb, "Quadrants", "Where students sit — before and after",
                "Quadrant = literacy × readiness. Band = the overall AI profile.")
    headers = ["Department", "Measure", "Group", "Before", "Before %", "After", "After %"]
    rows = []
    for o in every:
        for kind, mix in (("Quadrant", o["quadrants"]), ("Band", o["bands"])):
            for x in mix:
                rows.append([o["dept"], kind, x["label"], x["before"], x["before_pct"],
                             x["after"], x["after_pct"]])
    _table(ws, headers, rows, [42, 11, 24, 10, 10, 10, 10])

    # ── PRaiSE ──────────────────────────────────────────────────────────────
    ws = _sheet(wb, "PRaiSE", "Who wants to contribute to PRaiSE, by pillar",
                "From the post survey: the one pillar each student chose. Students who chose "
                "None are counted on the Summary sheet, not listed here.")
    headers = ["Department", "Pillar", "Student", "Email"]
    rows = [[o["dept"], p["pillar"], st["name"], st["email"]]
            for o in departments for p in o["praise"]["pillars"] for st in p["students"]]
    widths = [42, 20, 30, 34]
    if not include_emails:
        headers, rows, widths = _drop(headers, rows, widths, "Email")
    _table(ws, headers, rows, widths)

    # ── Entrepreneurs ───────────────────────────────────────────────────────
    ws = _sheet(wb, "Entrepreneurs", "Students from an entrepreneur family",
                "From the post survey: a parent whose occupation is Entrepreneur, with the "
                "business they named.")
    headers = ["Department", "Student", "Email", "Parent", "Parent's name", "Business",
               "Business type"]
    rows = [[o["dept"], st["name"], st["email"], par["parent"], par["name"],
             par["business"], par["type"]]
            for o in departments for st in o["entrepreneurs"]["students"]
            for par in st["parents"]]
    widths = [42, 30, 34, 10, 26, 30, 22]
    if not include_emails:
        headers, rows, widths = _drop(headers, rows, widths, "Email")
    _table(ws, headers, rows, widths)

    # ── Students ────────────────────────────────────────────────────────────
    ws = _sheet(wb, "Students", "Every student — before, after and the change",
                "Blank scores mean that survey was not filled.")
    headers = ["Department", "Student", "Email", "Campus", "Baseline filled", "Post filled",
               "Literacy before", "Readiness before", "Overall before", "Quadrant before",
               "Literacy after", "Readiness after", "Overall after", "Quadrant after",
               "Change", "PRaiSE pillar", "Entrepreneur family"]
    rows = []
    for o in departments:
        for st in o["students"]:
            b, a = st["before"], st["after"]
            rows.append([o["dept"], st["name"], st["email"], st["campus"],
                         "Yes" if st["baseline_done"] else "No",
                         "Yes" if st["post_done"] else "No",
                         b["lit"], b["read"], b["overall"], b["quadrant"],
                         a["lit"], a["read"], a["overall"], a["quadrant"],
                         st["change"], st["praise"],
                         "Yes" if st["entrepreneur_family"] else "No"])
    widths = [42, 30, 34, 12, 11, 11, 11, 11, 11, 18, 11, 11, 11, 18, 10, 18, 12]
    if not include_emails:
        headers, rows, widths = _drop(headers, rows, widths, "Email")
    _table(ws, headers, rows, widths)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def workbook_filename(scope: str) -> str:
    name = f"AI_Survey_Outcome_Impact_{scope}.xlsx"
    return "".join(c if (c.isalnum() or c in "._-") else "_" for c in name)


# Section titles re-exported for anyone formatting the sheet names elsewhere.
SECTION_NAMES = {k: SECTION_TITLES.get(k, k) for k in SECTIONS}
